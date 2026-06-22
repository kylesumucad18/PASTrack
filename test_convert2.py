import requests
import base64
import os
import sys

try:
    from docx import Document
    document = Document()
    document.add_heading('Test Document', 0)
    document.add_paragraph('This is a test document.')
    document.save('test_valid.docx')
except ImportError:
    print("python-docx not installed, cannot create a valid docx.")
    sys.exit(1)

with open('test_valid.docx', 'rb') as f:
    file_content = f.read()

api_secret = "OoLm1SeWltZvzUzwinJQfl3gyMA6nCAz"
url = f"https://v2.convertapi.com/convert/docx/to/pdf?Secret={api_secret}"

files = {
    'File': ("test_valid.docx", file_content)
}

response = requests.post(url, files=files)
print("Status:", response.status_code)
print("Response:", response.text)

if response.status_code == 200:
    data = response.json()
    b64 = data['Files'][0]['FileData']
    with open('test_valid.pdf', 'wb') as f:
        f.write(base64.b64decode(b64))
    print("PDF saved.")
